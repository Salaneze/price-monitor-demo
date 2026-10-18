"""python -m unittest  (pytest тоже подхватит)"""
import sqlite3
import tempfile
import unittest
from dataclasses import replace
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from monitor import export, store
from monitor.notify import CHUNK, format_changes
from monitor.scraper import Product, parse_category_page, parse_price, parse_product_page, parse_qty

FIX = Path(__file__).parent / "fixtures"
BASE = "https://books.toscrape.com/catalogue/category/books/travel_2/index.html"


def product(upc="u1", price="10.00", qty=5, category="Travel", **kw) -> Product:
    return Product(upc=upc, title=kw.get("title", f"Книга {upc}"), url=kw.get("url", f"https://x/{upc}"),
                   category=category, price=Decimal(price), qty=qty, rating=3)


class ParseTest(unittest.TestCase):
    def test_category_page(self):
        links, nxt = parse_category_page((FIX / "cat.html").read_text("utf-8"), BASE)
        self.assertEqual(len(links), 11)
        self.assertTrue(links[0].startswith("https://books.toscrape.com/catalogue/"))
        self.assertIsNone(nxt)

    def test_pagination(self):
        url = "https://books.toscrape.com/catalogue/category/books/mystery_3/index.html"
        links, nxt = parse_category_page((FIX / "cat_paged.html").read_text("utf-8"), url)
        self.assertEqual(len(links), 20)
        self.assertEqual(nxt, "https://books.toscrape.com/catalogue/category/books/mystery_3/page-2.html")

    def test_product_page(self):
        p = parse_product_page((FIX / "prod.html").read_text("utf-8"), "https://x/p")
        self.assertEqual(p.upc, "a22124811bfa8350")
        self.assertEqual(p.price, Decimal("45.17"))
        self.assertEqual(p.qty, 19)
        self.assertEqual(p.category, "Travel")
        self.assertEqual(p.rating, 2)

    def test_dirty_values(self):
        self.assertEqual(parse_price("Â£51.77"), Decimal("51.77"))
        self.assertEqual(parse_price("1 299,50 ₽".replace(" ", "")), Decimal("1299.50"))
        self.assertEqual(parse_qty("Out of stock"), 0)
        with self.assertRaises(ValueError):
            parse_price("нет в наличии")


class SnapshotTest(unittest.TestCase):
    def setUp(self):
        self.db = sqlite3.connect(":memory:")
        self.db.row_factory = sqlite3.Row
        self.db.executescript(store.SCHEMA)

    def kinds(self, changes):
        return sorted((c.kind, c.upc) for c in changes)

    def test_first_run_is_silent(self):
        self.assertEqual(store.apply_snapshot(self.db, [product("a"), product("b")]), [])

    def test_rerun_updates_not_duplicates(self):
        store.apply_snapshot(self.db, [product("a")])
        store.apply_snapshot(self.db, [product("a", url="https://x/moved")])
        rows = self.db.execute("SELECT url FROM products").fetchall()
        self.assertEqual([r["url"] for r in rows], ["https://x/moved"])
        # цена не менялась — второй строки в истории нет
        self.assertEqual(self.db.execute("SELECT count(*) FROM history").fetchone()[0], 1)

    def test_price_and_stock_changes(self):
        store.apply_snapshot(self.db, [product("a"), product("b"), product("c", qty=0), product("d")])
        changes = store.apply_snapshot(self.db, [
            product("a", price="8.50"),
            product("b", price="12.00", qty=0),
            product("c", qty=3),
            product("d"),
            product("e"),
        ])
        self.assertEqual(self.kinds(changes), [
            ("back", "c"), ("cheaper", "a"), ("new", "e"), ("out", "b"), ("pricier", "b"),
        ])
        cheaper = next(c for c in changes if c.kind == "cheaper")
        self.assertEqual((cheaper.old_price, cheaper.new_price), (Decimal("10.00"), Decimal("8.50")))

    def test_gone_only_when_really_gone(self):
        store.apply_snapshot(self.db, [product("a"), product("b"), product("x", category="Poetry")])
        # b упал по сети, x из категории, которую сейчас не обходили
        changes = store.apply_snapshot(self.db, [product("a")], failed_urls=["https://x/b"])
        self.assertEqual(changes, [])
        changes = store.apply_snapshot(self.db, [product("a")])
        self.assertEqual(self.kinds(changes), [("gone", "b")])
        # повторно о том же не кричим
        self.assertEqual(store.apply_snapshot(self.db, [product("a")]), [])
        # вернулся — сообщаем
        self.assertEqual(self.kinds(store.apply_snapshot(self.db, [product("a"), product("b")])), [("back", "b")])

    def test_money_is_exact(self):
        store.apply_snapshot(self.db, [product("a", price="51.77")])
        self.assertEqual(store.apply_snapshot(self.db, [product("a", price="51.77")]), [])

    def test_xlsx(self):
        store.apply_snapshot(self.db, [product("a"), product("b", category="Poetry")])
        store.apply_snapshot(self.db, [product("a", price="9.99"), product("b", category="Poetry")])
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "out.xlsx"
            export.to_xlsx(self.db, path)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Товары", "История цен"])
            self.assertEqual(wb["Товары"].max_row, 3)
            self.assertEqual(wb["История цен"].max_row, 4)
            self.assertEqual(wb["История цен"]["D2"].value, 9.99)


class NotifyTest(unittest.TestCase):
    def test_long_list_is_split(self):
        changes = [store.Change("cheaper", str(i), "Очень длинное название книги " * 3, f"https://x/{i}",
                                Decimal("10"), Decimal("9")) for i in range(100)]
        msgs = format_changes(changes)
        self.assertGreater(len(msgs), 1)
        self.assertTrue(all(len(m) <= CHUNK + 300 for m in msgs))
        self.assertEqual(sum(m.count("📉") for m in msgs), 100)
        self.assertIn("£10 → £9", msgs[0])


if __name__ == "__main__":
    unittest.main()

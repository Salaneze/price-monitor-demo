"""Выгрузка в таблицы: Excel всегда, Google Sheets — если заданы ключи."""
import logging
import sqlite3
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

PRODUCT_HEADER = ["Артикул", "Название", "Категория", "Цена", "Остаток", "Рейтинг",
                  "В продаже", "Впервые замечен", "Последний обход", "Ссылка"]
HISTORY_HEADER = ["Артикул", "Название", "Когда", "Цена", "Остаток"]


def _rows(db: sqlite3.Connection) -> tuple[list[list], list[list]]:
    products = [
        [r["upc"], r["title"], r["category"], float(Decimal(r["price"])), r["qty"], r["rating"],
         "да" if r["active"] else "нет", r["first_seen"], r["last_seen"], r["url"]]
        for r in db.execute("SELECT * FROM products ORDER BY category, title")
    ]
    history = [
        [r["upc"], r["title"], r["ts"], float(Decimal(r["price"])), r["qty"]]
        for r in db.execute(
            """SELECT h.upc, p.title, h.ts, h.price, h.qty
               FROM history h JOIN products p USING (upc)
               ORDER BY h.ts DESC, h.rowid DESC"""
        )
    ]
    return products, history


def to_xlsx(db: sqlite3.Connection, path: Path) -> None:
    products, history = _rows(db)
    wb = Workbook()
    for ws, header, rows in (
        (wb.active, PRODUCT_HEADER, products),
        (wb.create_sheet(), HISTORY_HEADER, history),
    ):
        ws.append(header)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(header, 1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    wb.worksheets[0].title = "Товары"
    wb.worksheets[1].title = "История цен"
    path.parent.mkdir(parents=True, exist_ok=True)
    # Открытый в Excel файл на Windows заблокирован — пишем рядом, а не падаем
    try:
        wb.save(path)
    except PermissionError:
        alt = path.with_stem(path.stem + "_new")
        wb.save(alt)
        log.warning("%s открыт в другой программе, сохранил в %s", path, alt)


def to_gsheet(db: sqlite3.Connection, sheet_id: str, credentials: str) -> None:
    # gspread нужен только тем, кто выгружает в Google — не тянем его в обязательные
    import gspread

    products, history = _rows(db)
    book = gspread.service_account(filename=credentials).open_by_key(sheet_id)
    for title, header, rows in (
        ("Товары", PRODUCT_HEADER, products),
        ("История цен", HISTORY_HEADER, history),
    ):
        try:
            ws = book.worksheet(title)
        except gspread.WorksheetNotFound:
            ws = book.add_worksheet(title, rows=len(rows) + 10, cols=len(header))
        ws.clear()
        ws.update([header] + rows, value_input_option="USER_ENTERED")
        ws.freeze(rows=1)
